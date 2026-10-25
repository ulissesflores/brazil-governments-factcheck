"""Constroi series/*.csv (CSV longo) e RESULTADO.md a partir do bruto/ hasheado. Sem rede. Nada estimado, nada emendado.

Convencao de data do SGS trimestral: a data e o PRIMEIRO mes do trimestre e o valor e a posicao de FIM de trimestre
(conferido: SGS 3684 em 01/04/2026 = 407.102,75 = jun/2026 na Tabela 15 da Nota de set/2026 e na DivExt_T.xls).
Aqui o campo mes_ou_trimestre recebe o mes de FIM do trimestre (3, 6, 9, 12) para que dezembro seja dezembro.
"""
import csv, json
from pathlib import Path

P = Path(__file__).resolve().parent.parent
man = {}
for l in (P / "manifest.jsonl").read_text().splitlines():
    m = json.loads(l); man[Path(m["arquivo"]).name] = m


def vint(fname):
    m = man[fname]
    return f"bruto/{fname} adquirido {m['timestamp_utc']} sha256 {m['sha256'][:12]}"


def load(c):
    d = json.load(open(P / f"bruto/sgs_{c}.json"))
    out = {}
    for x in d:
        ano, mes = int(x["data"][6:]), int(x["data"][3:5])
        assert (ano, mes) not in out, (c, ano, mes)
        out[(ano, mes)] = float(x["valor"])
    return out


TRIM = "trimestral; data do SGS = primeiro mes do trimestre, posicao de fim de trimestre; aqui mes = fim do trimestre"
DESC = "DESCONTINUADA pelo BCB em fev/2026 (Nota Explicativa de 24/02/2026); publicada como esta, sem prolongamento"
# codigo: (nome, definicao_linha, unidade, nota, trimestral?)
S = {
 3684: ("divida_externa_bruta_restrita_usd_mi", "Divida externa bruta, conceito restrito do BCB ('a chamada divida externa bruta': sem operacoes intercompanhia e sem titulos negociados no mercado domestico), posicao de fim de trimestre", "US$ milhoes", TRIM, True),
 25563: ("divida_externa_bruta_bpm6_usd_mi", "Divida externa bruta inclusive operacoes intercompanhia e titulos de divida negociados no mercado domestico em poder de nao residentes (conceito BPM6 / EDS 2013, 'a divida externa brasileira' segundo a Nota Explicativa)", "US$ milhoes", TRIM, True),
 7941: ("intercompanhia_usd_mi", "Operacoes intercompanhia (passivo de investimento direto), posicao de fim de trimestre", "US$ milhoes", TRIM, True),
 22160: ("titulos_mercado_domestico_nao_residentes_usd_mi", "Titulos de divida negociados no mercado domestico em poder de nao residentes, posicao de fim de trimestre", "US$ milhoes", TRIM, True),
 21521: ("divida_externa_setor_publico_restrita_usd_mi", "Divida externa bruta do setor publico, conceito restrito (governo geral, Banco Central e entidades publicas), posicao de fim de trimestre", "US$ milhoes", TRIM + "; o BCB revisou a divisao publico/privado de dez/2020 a set/2025 na Nota de mar/2026", True),
 24010: ("pii_liquida_usd_mi", "Posicao de investimento internacional liquida (ativos externos menos passivos externos), fim de trimestre", "US$ milhoes", TRIM, True),
 11407: ("divida_externa_restrita_pib_pct", "Divida externa bruta (conceito restrito, SGS 3684) / PIB acumulado em 12 meses em US$ (SGS 4192), %", "Percentual do PIB", DESC + "; " + TRIM, True),
 3685: ("divida_externa_liquida_usd_mi", "Divida externa liquida, conceito antigo do BCB (coluna M da Tabela 31 das Notas ate jan/2026): divida externa bruta restrita menos reservas internacionais, creditos brasileiros no exterior e haveres de bancos comerciais no exterior; negativo = credor externo liquido", "US$ milhoes", TRIM + "; AINDA PUBLICADA pelo BCB (conferido 2026-10-24); coincide com a Tabela 31 da Nota de jan/2026 nos trimestres comuns (jun/2024-set/2025); dez/2025 revisado pelo BCB de -17.217,78 (estimado na Nota) para -8.701,59", True),
 11409: ("divida_externa_liquida_pib_pct", "Divida externa liquida / PIB, % (conceito antigo do BCB: divida externa bruta restrita menos reservas internacionais, creditos brasileiros no exterior e haveres de bancos comerciais; negativo = credor externo liquido)", "Percentual do PIB", DESC + " porque a construcao original nao adere ao BPM6/EDS 2013; " + TRIM, True),
 11408: ("divida_externa_setor_publico_pct", "Divida externa bruta do setor publico (SGS 21521) / divida externa bruta restrita (SGS 3684), %", "Percentual", DESC + "; " + TRIM, True),
 12506: ("pii_liquida_pib_pct", "Posicao de investimento internacional liquida (SGS 24010) / PIB acumulado em 12 meses em US$ (SGS 4192), %", "Percentual do PIB", DESC + "; " + TRIM, True),
 4192: ("pib_usd_mi_mensal", "PIB mensal em US$ milhoes correntes (denominador das razoes do BCB, somado em 12 meses)", "US$ milhoes", "mensal", False),
 22705: ("exportacoes_bens_servicos_usd_mi", "Exportacoes de bens e servicos, mensal", "US$ milhoes", "mensal", False),
 5760: ("nfsp_juros_nominais_12m_pib_pct", "NFSP sem desvalorizacao cambial, juros nominais, setor publico consolidado, fluxo acumulado em 12 meses, % do PIB", "Percentual do PIB", "mensal; nome do catalogo CKAN do BCB (dataset 5760)", False),
 5727: ("nfsp_resultado_nominal_12m_pib_pct", "NFSP sem desvalorizacao cambial, resultado nominal, setor publico consolidado, fluxo acumulado em 12 meses, % do PIB (positivo = deficit, convencao do BCB)", "Percentual do PIB", "mensal; nome do catalogo CKAN do BCB (dataset 5727); sinal do BCB mantido aqui", False),
 5793: ("nfsp_resultado_primario_12m_pib_pct", "NFSP sem desvalorizacao cambial, resultado primario, setor publico consolidado, fluxo acumulado em 12 meses, % do PIB (positivo = deficit, convencao do BCB)", "Percentual do PIB", "mensal; nome do catalogo CKAN do BCB (dataset 5793); sinal do BCB mantido aqui", False),
}
CAB = "indicador,fonte,definicao_linha,ano,mes_ou_trimestre,valor,unidade,nota,vintage".split(",")
(P / "series").mkdir(exist_ok=True)
data = {}
for c, (nome, defin, un, nota, trim) in S.items():
    data[c] = load(c)
    with open(P / "series" / f"{nome}_sgs{c}.csv", "w", newline="") as f:
        w = csv.writer(f); w.writerow(CAB)
        for (a, m), v in sorted(data[c].items()):
            if trim:
                assert m in (1, 4, 7, 10), (c, a, m)
                m += 2
            w.writerow([nome, f"BCB SGS {c}", defin, a, m, v, un, nota, vint(f"sgs_{c}.json")])

# Tabela 31 da Nota de jan/2026: ultimos valores OFICIAIS da divida externa liquida (conceito antigo), em US$ milhoes
import openpyxl  # noqa: E402
wb = openpyxl.load_workbook(P / "bruto/202601_Tabelas_de_estatisticas_do_setor_externo.xlsx", read_only=True, data_only=True)
tab = next(ws for ws in wb.worksheets if any("sustentabilidade" in str(c).lower() for row in ws.iter_rows(min_row=1, max_row=2, values_only=True) for c in row if c))
rows = list(tab.iter_rows(values_only=True))
anos, meses = rows[3], rows[5]
colunas = []
ano = None
for i in range(1, len(meses)):
    if anos[i] not in (None, ""):
        ano = int(float(anos[i]))
    if meses[i]:
        mes = {"Jan": 1, "Fev": 2, "Mar": 3, "Abr": 4, "Mai": 5, "Jun": 6, "Jul": 7, "Ago": 8, "Set": 9, "Out": 10, "Nov": 11, "Dez": 12}[str(meses[i]).split()[0]]
        colunas.append((i, ano, mes, "(e)" in str(meses[i])))
LINHAS31 = {
    "Dívida externa bruta3/": "tab31_divida_externa_bruta_usd_mi",
    "Reservas internacionais (J)": "tab31_reservas_usd_mi",
    "Créditos brasileiros no ext": "tab31_creditos_brasileiros_exterior_usd_mi",
    "Haveres de bancos comerciais": "tab31_haveres_bancos_comerciais_usd_mi",
    "Dívida externa líquida (M)": "tab31_divida_externa_liquida_usd_mi",
    "Dívida externa líquida / PIB": "tab31_divida_externa_liquida_pib_pct",
    "Dívida externa bruta / PIB": "tab31_divida_externa_bruta_pib_pct",
}
v31 = vint("202601_Tabelas_de_estatisticas_do_setor_externo.xlsx")
with open(P / "series" / "tabela31_nota_jan2026.csv", "w", newline="") as f:
    w = csv.writer(f); w.writerow(CAB)
    for row in rows:
        lab = str(row[0] or "").strip()
        for prefixo, nome in LINHAS31.items():
            if lab.startswith(prefixo):
                un = "Percentual" if "pct" in nome else "US$ milhoes"
                for i, a, m, est in colunas:
                    if isinstance(row[i], (int, float)):
                        w.writerow([nome, "BCB, Nota para a Imprensa - Setor Externo, jan/2026, Tabela 31", lab, a, m, round(float(row[i]), 2), un, "(e) estimado pelo BCB" if est else "", v31])
                break

# RESULTADO.md
def ult(c):
    k = max(data[c]); m = k[1] + 2 if S[c][4] else k[1]
    return f"{data[c][k]:,.1f} ({k[0]}-{m:02d})".replace(",", "X").replace(".", ",").replace("X", ".")
def em(c, a, m):
    v = data[c].get((a, m)); return "sem dado" if v is None else f"{v:,.1f}".replace(",", "X").replace(".", ",").replace("X", ".")
md = ["# RESULTADO - coleta externo-fase3 (divida externa e custo da divida publica), 2026-10-23 (SGS 3685 acrescentada e errata em 2026-10-24)", "",
      "Pasta: `fontes/dados/externo-fase3/`. Cadeia: `manifest.jsonl` (todo download por `govfc.chain add`). Scripts: `scripts/baixar.py` (rede) e `scripts/construir.py` (sem rede).",
      "Dimensao `ext` do pre-registro original (PREREGISTRATION.md, 'External position: ... external debt'): camada descritiva da fase 1, sem emenda. Codigos do SGS nomeados pela Nota Explicativa do BCB de 24/02/2026 (`bruto/202602_Texto_de_estatisticas_do_setor_externo.pdf`).", "",
      "## Os dois conceitos de divida externa bruta do BCB (US$ milhoes, fim de trimestre)", "",
      "| Serie | Inicio | dez/2002 | dez/2010 | dez/2016 | dez/2022 | dez/2025 | Ultimo |", "|---|---|---|---|---|---|---|---|"]
for c in (3684, 25563, 21521, 7941, 22160):
    k0 = min(data[c]); md.append(f"| {S[c][0]} (SGS {c}) | {k0[0]}-{k0[1]+2:02d} | {em(c,2002,10)} | {em(c,2010,10)} | {em(c,2016,10)} | {em(c,2022,10)} | {em(c,2025,10)} | {ult(c)} |")
md += ["", "## Razoes publicadas pelo BCB (%, descontinuadas em fev/2026; publicadas como estao)", "",
       "| Serie | Inicio | 1984T4 | 1999T4 | 2002T4 | 2007T4 | 2010T4 | 2016T4 | 2022T4 | Ultimo |", "|---|---|---|---|---|---|---|---|---|---|"]
for c in (11407, 11409, 11408, 12506):
    k0 = min(data[c]); md.append(f"| {S[c][0]} (SGS {c}) | {k0[0]}-{k0[1]+2:02d} | {em(c,1984,10)} | {em(c,1999,10)} | {em(c,2002,10)} | {em(c,2007,10)} | {em(c,2010,10)} | {em(c,2016,10)} | {em(c,2022,10)} | {ult(c)} |")
md += ["", "## Custo da divida publica (NFSP, 12 meses, % do PIB, convencao do BCB: positivo = deficit)", "",
       "| Serie | Inicio | dez/2003 | dez/2010 | dez/2015 | dez/2016 | dez/2019 | dez/2022 | Ultimo |", "|---|---|---|---|---|---|---|---|---|"]
for c in (5760, 5727, 5793):
    k0 = min(data[c]); md.append(f"| {S[c][0]} (SGS {c}) | {k0[0]}-{k0[1]:02d} | {em(c,2003,12)} | {em(c,2010,12)} | {em(c,2015,12)} | {em(c,2016,12)} | {em(c,2019,12)} | {em(c,2022,12)} | {ult(c)} |")
md += ["", "## Divida externa liquida em US$ (SGS 3685, AINDA PUBLICADA; a Tabela 31 foi descontinuada em fev/2026)", "",
       "Definicao do BCB (coluna M da Tabela 31): liquida = bruta (restrita) - reservas internacionais - creditos brasileiros no exterior - haveres de bancos comerciais. Negativo = credor externo liquido. Tabela 31 da Nota de jan/2026 em `series/tabela31_nota_jan2026.csv`; a serie SGS 3685 coincide com ela em todos os trimestres comuns (jun/2024-set/2025) e traz dez/2025 revisado.", "",
       "| Serie | Inicio | mar/1990 | dez/2002 | set/2007 | dez/2007 | dez/2010 | dez/2016 | dez/2022 | dez/2024 | dez/2025 | Ultimo |", "|---|---|---|---|---|---|---|---|---|---|---|---|"]
for c in (3685,):
    k0 = min(data[c]); md.append(f"| {S[c][0]} (SGS {c}) | {k0[0]}-{k0[1]+2:02d} | {em(c,1990,1)} | {em(c,2002,10)} | {em(c,2007,7)} | {em(c,2007,10)} | {em(c,2010,10)} | {em(c,2016,10)} | {em(c,2022,10)} | {em(c,2024,10)} | {em(c,2025,10)} | {ult(c)} |")
md += ["",
       "ERRATA (2026-10-24). A primeira versao deste RESULTADO (2026-10-23) dizia que a noticia que motivou o pedido (BP Money, 21/10/2026; `bruto/noticia_bpmoney_2026-10-21.html`) calculava 'bruta - reservas' e que nao existia valor oficial de 2026 em nenhum conceito. As duas afirmacoes estavam erradas: o nivel da divida externa liquida continua publicado pelo BCB (SGS 3685) e e exatamente o numero da manchete (+5.083,49 em jun/2026, primeiro valor positivo desde set/2007); os 'US$ 402 bi de reservas' da manchete sao as reservas mais os creditos brasileiros no exterior e os haveres de bancos, o conjunto de ativos que o BCB subtrai (a propria noticia diz que os creditos ja estao contabilizados). O que foi descontinuado em fev/2026 e a RAZAO sobre o PIB (SGS 11409) e a Tabela 31. Nada foi calculado nem prolongado aqui.", "",
       "## Leitura", "",
       "'Sem dado' = a serie nao existe naquele ponto; nada estimado nem interpolado. As duas medidas de divida bruta nunca sao somadas nem emendadas: 25563 = 3684 + 7941 + 22160 por construcao do BCB, e a pagina mostra cada uma com o proprio nome. O WDI/IDS (DT.DOD.DECT.CD, GN.ZS) e um TERCEIRO conceito (Banco Mundial), usado so na comparacao com vizinhos na mesma fonte (colecao `mundo`)."]
(P / "RESULTADO.md").write_text("\n".join(md) + "\n", encoding="utf-8")
print("series:", len(list((P / "series").glob("*.csv"))))
